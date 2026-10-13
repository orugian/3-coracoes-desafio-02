#!/usr/bin/env python3
"""Reproduz os números citados em docs/arquitetura/ a partir das fontes do desafio.

Uso (na raiz do repositório):
    python docs/arquitetura/verificacao/verificar_numeros.py

Requer: Python 3.10+ e openpyxl. Só lê arquivos; não grava nada.
Fontes: pre-solucao/pedidos-sem-fronteiras/*.xlsx|xlsm e dados/*.json.
Saída: relatório por seção e um resumo "N de M verificações conferem".

O que este script prova: que cada número citado na documentação sai das fontes
do desafio com as regras descritas no texto (reprodutibilidade). Ele é uma
reimplementação em Python dessas regras, não uma cópia do código do piloto.
O que ele NÃO prova: que as regras estão certas para a 3corações. Posições de
pallet (28/20/10) vêm do artefato do time, confirmadas por ele como dado correto,
e a altura interna dos veículos é premissa (ver ANEXO_EVIDENCIAS.md). As caixas por
pallet de 12034109 (98) e 20908885 (77) são decisão do time por heurística (Q12); o
script confere que elas respeitam a regra declarada, não que a logística as confirma.
O motor do piloto (commit 636ac8d) foi executado à parte e dá os mesmos números.
"""
import glob
import io
import json
import math
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "pre-solucao" / "pedidos-sem-fronteiras"
DADOS = ROOT / "dados"
EPS = 1e-9

CHECKS = []  # (descrição, esperado, obtido, ok)


def check(desc, expected, got, tol=0.006):
    if isinstance(expected, (int, float)) and isinstance(got, (int, float)):
        ok = abs(expected - got) <= tol
    else:
        ok = expected == got
    CHECKS.append((desc, expected, got, ok))
    print(f"  [{'ok' if ok else 'DIVERGE'}] {desc}: esperado {expected} | obtido {got}")


def section(title):
    print(f"\n== {title}")


# ---------------------------------------------------------------- leitura
def first(pattern):
    return glob.glob(str(SRC / pattern))[0]


conf = openpyxl.load_workbook(first("Confer*.xlsx"), data_only=True).active
mats = openpyxl.load_workbook(first("Materiais*.xlsx"), data_only=True).active
regs = openpyxl.load_workbook(first("Registros*.xlsx"), data_only=True).active
ob = openpyxl.load_workbook(SRC / "ORDEN BASIS TESTE.xlsm", data_only=True, keep_vba=True, read_only=True)

CX, ROWS = {}, {}
for r in conf.iter_rows(min_row=2, values_only=True):
    m, um, cnt, net, gross, unit = str(r[2]), r[5], r[6], r[15], r[16], r[17]
    k = 0.001 if unit == "G" else 1.0
    ROWS.setdefault(m, {})[um] = dict(cnt=cnt, net=net * k, gross=gross * k, gross_raw=gross, net_raw=net, unit=unit)
    if um == "CX":
        CX[m] = dict(L=r[9], W=r[10], H=r[11], vol=r[13], un=cnt, net=net * k, gross=gross * k, unit=unit,
                     gross_raw=gross, net_raw=net, shelf=r[20])
LAYERS = {str(r[0]): (int(r[5]), int(r[6])) for r in mats.iter_rows(min_row=2, values_only=True)}  # lastro, altura
CXPAL = {m: a * b for m, (a, b) in LAYERS.items()}
SAP_PRICE = {str(r[8]): r[16] for r in regs.iter_rows(min_row=2, values_only=True)}
SKUS = sorted(CX)

TAB = {str(r[5]): r[8] for r in ob["TAB"].iter_rows(min_row=3, values_only=True) if r[5]}  # preço usado pela OB
TAB_CXPAL = {str(r[5]): r[12] for r in ob["TAB"].iter_rows(min_row=3, values_only=True) if r[5]}  # cx/pallet da aba TAB
PALET = [(r[1], str(r[2]), r[3], r[7]) for r in ob["PALETIZAÇÃO"].iter_rows(min_row=3, values_only=True) if r[2]]
PALET_UN = [(r[1], str(r[2]), r[6]) for r in ob["PALETIZAÇÃO"].iter_rows(min_row=3, values_only=True) if r[2]]  # un/pallet
AUXLANG = {}
for r in ob["AUX"].iter_rows(min_row=1, max_row=60, values_only=True):
    lang, mat = r[2], r[3]
    if lang in ("PORTUGUÊS", "ESPAÑOL", "ENGLISH") and isinstance(mat, (int, float)):
        AUXLANG.setdefault(str(int(mat)), set()).add(lang)
ORDEM_AA2 = None
for r in ob["ORDEM"].iter_rows(min_row=2, max_row=2, values_only=True):
    ORDEM_AA2 = r[26]  # AA2 = X2/26

catalogo = json.load(io.open(DADOS / "catalogo.json", encoding="utf-8"))
internos = json.load(io.open(DADOS / "produtos_internos.json", encoding="utf-8"))
pedido = json.load(io.open(DADOS / "pedido.json", encoding="utf-8"))
container = json.load(io.open(DADOS / "container.json", encoding="utf-8"))["modais_logisticos"]
CAT = {str(i["material"]): i for i in catalogo["itens"]}
INT = {str(i["material"]): i for i in internos["itens"]}
VEH = {v["id_veiculo"]: v["capacidade_maxima"] for v in container}
# posições de pallet: dado do artefato (piloto 3C Conecta), confirmado pelo time (dado mockado, mas correto)
POS = {"carreta_padrao": 28, "container_40": 20, "container_20": 10}

# ---------------------------------------------------------------- motor mínimo
Q = {str(i["material"]): i["caixas"] for i in pedido["itens"]}


def line_totals(q):
    t = dict(gross=0.0, net=0.0, m3_decl=0.0, m3_dims=0.0, pallets=0.0, pos=0)
    for m, cx in q.items():
        c = CX[m]
        t["gross"] += cx * c["gross"]
        t["net"] += cx * c["net"]
        t["m3_decl"] += cx * c["vol"] / 1000
        t["m3_dims"] += cx * c["L"] * c["W"] * c["H"] / 1e9
        p = cx / CXPAL[m]
        t["pallets"] += p
        t["pos"] += math.ceil(p - EPS)
    return t


def occupancy(t, veh, pos_limit, mixed=False, tara=0.0):
    cap = VEH[veh]
    pos = math.ceil(t["pallets"] - EPS) if mixed else t["pos"]
    kg = t["gross"] + tara * pos
    return dict(pos=pos, pos_pct=pos / pos_limit * 100, kg_pct=kg / cap["peso_bruto_kg"] * 100,
                m3_pct=t["m3_decl"] / cap["cubagem_m3"] * 100)


# ================================================================ 1) pedido-teste
section("1) Pedido-teste (4 linhas, cliente 11131) recalculado com dados SAP")
T = line_totals(Q)
check("bruto das caixas (kg)", 21024.24, round(T["gross"], 2))
check("líquido (kg)", 19230.00, round(T["net"], 2))
check("cubagem declarada (m³)", 58.88, round(T["m3_decl"], 2))
check("cubagem pelas medidas C×L×A (m³)", 59.04, round(T["m3_dims"], 2))
check("pallets (soma fracionária, como a OB)", 28.886, round(T["pallets"], 3))
check("posições físicas (cada pallet incompleto ocupa uma)", 30, T["pos"])
o = occupancy(T, "carreta_padrao", POS["carreta_padrao"])
check("carreta, 28 posições (% de posições)", 107.1, round(o["pos_pct"], 1))
check("carreta, % de peso", 84.1, round(o["kg_pct"], 1))
check("carreta, % de cubagem", 81.8, round(o["m3_pct"], 1))
check("carreta, 26 posições (dica de AA2 na OB)", 115.4, round(30 / 26 * 100, 1))
check("carreta com pallet misto (posições)", 29, occupancy(T, "carreta_padrao", 28, mixed=True)["pos"])
check("carreta com pallet misto (%)", 103.6, round(occupancy(T, "carreta_padrao", 28, mixed=True)["pos_pct"], 1))
c40 = occupancy(T, "container_40", POS["container_40"])
check("contêiner 40', % de posições", 150.0, round(c40["pos_pct"], 1))
check("contêiner 40', % de cubagem", 96.8, round(c40["m3_pct"], 1))
check("contêiner 20', % de cubagem", 198.3, round(occupancy(T, "container_20", 10)["m3_pct"], 1))
tara = occupancy(T, "carreta_padrao", 28, tara=25.0)
check("carreta, % de peso com tara de 25 kg/pallet", 87.1, round(tara["kg_pct"], 1))
print(f"  só com cubagem e peso (dados do container.json) o pedido passa: peso {o['kg_pct']:.1f}%, cubagem {o['m3_pct']:.1f}%")

# faltas para fechar o próximo pallet e lote mínimo
for m, esperado in (("12034103", 24), ("12037417", 23), ("12144012", 118)):
    falta = math.ceil(CXPAL[m] * (math.floor(Q[m] / CXPAL[m] + EPS) + 1) - Q[m] - EPS)
    check(f"faltam caixas para fechar o pallet de {m}", esperado, falta)
check("12144012: pallets do pedido (lote 1 em LONT)", 1.44, round(Q["12144012"] / CXPAL["12144012"], 2))

# valores
price_motor = {"12031189": SAP_PRICE["12031189"], "12034103": SAP_PRICE["12034103"], "12144012": SAP_PRICE["12144012"],
               "12037417": TAB["12037417"]}  # 12037417 não tem preço no SAP: usa o do site da 3corações (Q6), igual ao da aba TAB
valor = sum(Q[m] * price_motor[m] for m in Q)
check("valor com preços SAP + 12037417 da OB (US$)", 192017.48, round(valor, 2))
check("valor sem o 12037417 (US$)", 184647.48, round(valor - Q["12037417"] * price_motor["12037417"], 2))
check("valor da OB (pedido.json) (US$)", 192005.84, pedido["totais"]["valor_total"])
check("diferença motor − OB (US$), vinda do preço do 12031189", 11.64, round(valor - pedido["totais"]["valor_total"], 2))

# ================================================================ 2) baseline da OB
section("2) O que a Order Basis calculou (células da aba ORDEM) × pedido.json × motor")
o_ = ob["ORDEM"]
cell = lambda a: o_[a].value
check("ORDEM!X2 peso bruto total da OB (kg)", 20151.04, round(cell("X2"), 2))
check("ORDEM!W2 peso líquido total da OB (kg)", 19070.0, cell("W2"))
check("ORDEM!Y2 valor total da OB (US$)", 192005.84, round(cell("Y2"), 2))
check("ORDEM!X4 cubagem total da OB (m³)", 58.88, round(cell("X4"), 2))
check("ORDEM!Y4 pallets da OB (soma de frações)", 28.886, round(cell("Y4"), 3))
check("ORDEM!W4 'EMBARQUE' calculado pela OB", "OK", cell("W4"))
check("pedido.json repete os totais das células da OB (bruto, líquido, valor)", True,
      abs(pedido["totais"]["peso_bruto_total"] - cell("X2")) < 0.01 and abs(pedido["totais"]["peso_liquido_total"] - cell("W2")) < 0.01
      and abs(pedido["totais"]["valor_total"] - cell("Y2")) < 0.01)
check("diferença de peso bruto OB × SAP (kg)", 873.20, round(T["gross"] - cell("X2"), 2))
check("diferença de peso bruto OB × SAP (%)", 4.3, round((T["gross"] / cell("X2") - 1) * 100, 1), tol=0.05)
check("diferença do valor: motor − OB (US$), vinda do preço do 12031189", 11.64, round(valor - cell("Y2"), 2))
f = next(i for i in pedido["itens"] if i["material"] == 12037417)
check("filtro 12037417: bruto por caixa na OB (kg)", 1.36, round(f["peso_bruto"] / f["caixas"], 2))
check("filtro 12037417: erro de peso da linha (%)", -41.9, round((f["peso_bruto"] / f["caixas"] / CX["12037417"]["gross"] - 1) * 100, 1), tol=0.05)
check("OB marcou 12144012 como não validado (lote de BHO)", False, next(i for i in pedido["itens"] if i["material"] == 12144012)["validado"])

# ================================================================ 2b) cadastro da caixa: aba PROD da OB × SAP
section("2b) Cadastro da caixa: aba PROD da Order Basis × SAP")
OBP = {}
for r in ob["PROD"].iter_rows(min_row=3, values_only=True):
    if r[7] == "CX" and str(r[4]) not in OBP:
        OBP[str(r[4])] = dict(L=r[11], W=r[12], H=r[13], vol=r[15], gross=r[17], net=r[19])
div_dims, div_weight = set(), set()
for m in SKUS:
    c, o_ = CX[m], OBP[m]
    if (c["L"], c["W"], c["H"]) != (o_["L"], o_["W"], o_["H"]) or abs(c["vol"] - o_["vol"]) > 0.001:
        div_dims.add(m)
    if abs(c["gross"] - o_["gross"]) > 0.0005 or abs(c["net"] - o_["net"]) > 0.0005:
        div_weight.add(m)
check("SKUs com medidas ou volume da caixa divergentes (OB × SAP)", ["12142501", "12144012", "12147801", "20908885"], sorted(div_dims))
check("SKUs com peso da caixa divergente (OB × SAP)", ["12034103", "12037417", "12147801", "20908885"], sorted(div_weight))
check("SKUs com cadastro da caixa divergente (união)", 6, len(div_dims | div_weight))

# ================================================================ 3) regras de qualidade
section("3) Regras de qualidade da importação, por SKU")
fire = {k: set() for k in ("VOL_MISMATCH", "LAYER_FILL_OUT_OF_RANGE", "PALLET_TOO_TALL", "MISSING_PRICE", "MIN_LOT_DEFAULTED",
                           "DUP_KEY", "GROSS_LT_NET", "UOM_G_TO_KG")}
MAX_H_MM = 2393  # premissa: altura interna de contêiner ISO; base do pallet 150 mm
for m in SKUS:
    c = CX[m]
    dims_dm3 = c["L"] * c["W"] * c["H"] / 1e6
    if abs(dims_dm3 - c["vol"]) / c["vol"] > 0.02:
        fire["VOL_MISMATCH"].add(m)
    fill = LAYERS[m][0] * c["L"] * c["W"] / 1.2e6
    if not 0.70 <= fill <= 1.02:
        fire["LAYER_FILL_OUT_OF_RANGE"].add(m)
    if LAYERS[m][1] * c["H"] + 150 > MAX_H_MM:
        fire["PALLET_TOO_TALL"].add(m)
    if m not in SAP_PRICE:
        fire["MISSING_PRICE"].add(m)
    if c["unit"] == "G":
        fire["UOM_G_TO_KG"].add(m)
        if c["gross"] < c["net"]:
            fire["GROSS_LT_NET"].add(m)
lont = {m for b, m, _, _ in PALET if b == "LONT"}
fire["MIN_LOT_DEFAULTED"] = {m for m in SKUS if m not in lont}
seen = {}
for b, m, d, _ in PALET:
    seen.setdefault((m, b), []).append(d)
fire["DUP_KEY"] = {m for (m, b), v in seen.items() if b == "LONT" and len(v) > 1}
expected_rules = {"VOL_MISMATCH": {"12144012"}, "LAYER_FILL_OUT_OF_RANGE": {"12034109"}, "PALLET_TOO_TALL": {"20908885"},
                  "MISSING_PRICE": {"12034109", "12037417", "20908885"}, "MIN_LOT_DEFAULTED": {"12034103"},
                  "DUP_KEY": {"12037417"}, "GROSS_LT_NET": {"12147801"}, "UOM_G_TO_KG": {"12147801"}}
for rule, exp in expected_rules.items():
    check(f"{rule} dispara em", sorted(exp), sorted(fire[rule]))
check("12147801: volume pelas medidas × declarado (%)", 1.4,
      round(abs(CX["12147801"]["L"] * CX["12147801"]["W"] * CX["12147801"]["H"] / 1e6 - CX["12147801"]["vol"]) / CX["12147801"]["vol"] * 100, 1), tol=0.05)
check("12144012: volume pelas medidas × declarado (%)", 7.0,
      round(abs(CX["12144012"]["L"] * CX["12144012"]["W"] * CX["12144012"]["H"] / 1e6 - CX["12144012"]["vol"]) / CX["12144012"]["vol"] * 100, 1), tol=0.05)
fill_pct = lambda m, lastro: lastro * CX[m]["L"] * CX[m]["W"] / 1.2e6 * 100  # % da base PBR 1,2 × 1,0 m
check("12034109: ocupação da base PBR com lastro 4 (%)", 28.4, round(fill_pct("12034109", 4), 1), tol=0.05)
check("12034109: ocupação da base PBR com lastro 14 (%)", 99.3, round(fill_pct("12034109", 14), 1), tol=0.05)
check("20908885: altura do pallet com base (m)", 2.65, round((LAYERS["20908885"][1] * CX["20908885"]["H"] + 150) / 1000, 2))
check("pallet padrão do café: altura com base (m)", 2.058, round((12 * 159 + 150) / 1000, 3))
check("duas linhas 12037417/LONT na PALETIZAÇÃO, lotes", [1, 2], sorted(l for b, m, d, l in PALET if m == "12037417" and b == "LONT"))
check("12144012: lote de BHO × LONT", [3.12, 1], [next(l for b, m, d, l in PALET if m == "12144012" and b == "BHO"),
                                                     next(l for b, m, d, l in PALET if m == "12144012" and b == "LONT")], tol=0.001)
check("filiais da PALETIZAÇÃO da OB", ["BHO", "LONT", "NAT"], sorted({b for b, _, _, _ in PALET}))
centros = {str(r[0]) for r in mats.iter_rows(min_row=2, values_only=True) if False} | {str(r[4]) for r in mats.iter_rows(min_row=2, values_only=True)}
check("centros presentes nos exports do SAP (só 2612; BHO e NAT não têm centro)", ["2612"], sorted(centros))

# ================================================================ 3b) Q12: caixas por pallet adotadas por heurística
section("3b) Caixas por pallet de 12034109 e 20908885 (Q12, heurística do time)")
# Decisão do time (Q12), registrada como master_override: lastro × camadas adotados.
# Regra: (1) limites rígidos (lastro até o limite por área; altura com base de 150 mm até 2.393 mm);
# (2) preferir o valor sustentado por mais de uma fonte que decomponha em inteiros; (3) conferir o
# preenchimento da base contra os demais SKUs. Não confirmada pela logística da 3corações.
ADOPTED = {"12034109": (14, 7), "20908885": (11, 7)}
LAY_A = {**LAYERS, **ADOPTED}
CXPAL_A = {m: a * b for m, (a, b) in LAY_A.items()}
area_limit = lambda m: math.floor(1.2e6 / (CX[m]["L"] * CX[m]["W"]))
height_mm = lambda m, cam: cam * CX[m]["H"] + 150
palet_cx = lambda m: sorted({round(u / CX[m]["un"], 2) for b, mm, u in PALET_UN if mm == m})
others_fill = [fill_pct(m, LAYERS[m][0]) for m in SKUS if m != "12034109"]
check("preenchimento da base dos outros 9 SKUs com o lastro do SAP (faixa, %)", (79.6, 97.4), (round(min(others_fill), 1), round(max(others_fill), 1)), tol=0.05)
check("limite de lastro por área da base: 12034109 e 20908885", (14, 11), (area_limit("12034109"), area_limit("20908885")))
check("12034109: cx/pallet nas abas TAB e PALETIZAÇÃO da OB (2.352 un ÷ 24, filiais BHO e LONT)", (98, [98.0]), (round(TAB_CXPAL["12034109"]), palet_cx("12034109")))
check("20908885: cx/pallet na PALETIZAÇÃO (924 un ÷ 12, filiais BHO, LONT e NAT) e na TAB", ([77.0], 60), (palet_cx("20908885"), round(TAB_CXPAL["20908885"])))
check("filiais com linha de PALETIZAÇÃO: 12034109 e 20908885", (["BHO", "LONT"], ["BHO", "LONT", "NAT"]),
      tuple(sorted({b for b, mm, _ in PALET_UN if mm == m}) for m in ("12034109", "20908885")))
check("20908885: o 60 da TAB ÷ lastro 11 do SAP (camadas, não inteiro)", 5.45, round(60 / LAYERS["20908885"][0], 2))
check("20908885: altura do pallet com 7 e com 10 camadas (mm)", (1900, 2650), (height_mm("20908885", 7), height_mm("20908885", 10)))
check("valores adotados (cx/pallet de 12034109 e 20908885)", (98, 77), (CXPAL_A["12034109"], CXPAL_A["20908885"]))
check("valores adotados: altura do pallet (mm, limite 2.393)", (1900, 1900), tuple(height_mm(m, LAY_A[m][1]) for m in ("12034109", "20908885")))
check("valores adotados: preenchimento da base (%)", (99.3, 93.0), tuple(round(fill_pct(m, LAY_A[m][0]), 1) for m in ("12034109", "20908885")), tol=0.05)
check("valores adotados: LAYER_FILL_OUT_OF_RANGE e PALLET_TOO_TALL disparam em", [],
      sorted(m for m in SKUS if not 70 <= fill_pct(m, LAY_A[m][0]) <= 102 or height_mm(m, LAY_A[m][1]) > MAX_H_MM))


def guillotine_max(W, H, a, b):
    """Máximo de caixas a×b (com rotação) num retângulo W×H só com cortes de guilhotina."""
    from functools import lru_cache
    lo, hi = min(a, b), max(a, b)
    raster = lambda lim: sorted({i * lo + j * hi for i in range(lim // lo + 1) for j in range((lim - i * lo) // hi + 1)})
    RX, RY = raster(W), raster(H)
    norm = lambda v, R: max(x for x in R if x <= v)

    @lru_cache(None)
    def f(w, h):
        best = 1 if (w >= a and h >= b) or (w >= b and h >= a) else 0
        for x in RX:
            if 0 < x < w and x <= w // 2 + hi:
                best = max(best, f(norm(x, RX), h) + f(norm(w - x, RX), h))
        for y in RY:
            if 0 < y < h and y <= h // 2 + hi:
                best = max(best, f(w, norm(y, RY)) + f(w, norm(h - y, RY)))
        return best
    return f(norm(W, RX), norm(H, RY))


check("guilhotina não serve de critério: máximo por camada no 12034109 e no 12031189 (lastro real 14)", (13, 12),
      tuple(guillotine_max(1200, 1000, int(CX[m]["L"]), int(CX[m]["W"])) for m in ("12034109", "12031189")))

# ================================================================ 4) JSON v2 × SAP
section("4) dados/*.json (v2) × SAP")
check("catálogo: itens", 10, len(CAT))
check("catálogo: preço == aba TAB da OB", 10, sum(abs(CAT[m]["preco_venda"] - TAB[m]) < 1e-6 for m in CAT))
check("catálogo: preço == Registros SAP (arredondado a 2 casas, 7 SKUs com preço)", 6,
      sum(m in SAP_PRICE and abs(CAT[m]["preco_venda"] - SAP_PRICE[m]) < 0.005 for m in CAT))
check("catálogo: preço diferente do SAP", [("12142146", 50.01, 50.27)], [(m, CAT[m]["preco_venda"], SAP_PRICE[m]) for m in CAT if m in SAP_PRICE and abs(CAT[m]["preco_venda"] - SAP_PRICE[m]) > 0.005])
check("catálogo: unidade_venda", ["CX"], sorted({c["unidade_venda"] for c in CAT.values()}))
mism = {}
for m in CAT:
    f = []
    if abs(CAT[m]["peso_liquido"] - CX[m]["net"]) > 1e-6: f.append("líquido")
    if abs(CAT[m]["peso_bruto"] - CX[m]["gross"]) > 1e-6: f.append("bruto")
    if abs(INT[m]["volume_max"] - CX[m]["vol"]) > 1e-3: f.append("volume")
    if abs(CAT[m]["paletizacao_caixas"] - CXPAL[m]) > 0.5: f.append("cx/pallet")
    if f: mism[m] = f
check("SKUs do JSON com valor ≠ SAP (peso, volume ou cx/pallet)", ["12034109", "12142501", "12144012", "12147801", "20908885"], sorted(mism))
check("peso líquido errado (SKUs)", ["12142501", "12144012", "12147801", "20908885"], sorted(m for m, f in mism.items() if "líquido" in f))
check("peso bruto errado (SKUs)", ["12147801"], sorted(m for m, f in mism.items() if "bruto" in f))
check("volume ≠ SAP (SKUs)", ["12142501", "20908885"], sorted(m for m, f in mism.items() if "volume" in f))
check("cx/pallet ≠ SAP lastro×altura (SKUs)", ["12034109", "20908885"], sorted(m for m, f in mism.items() if "cx/pallet" in f))
check("cx/pallet do JSON ≠ valor adotado na Q12 (SKUs; o 60 do 20908885 contra 77)", ["20908885"],
      sorted(m for m in CAT if abs(CAT[m]["paletizacao_caixas"] - CXPAL_A[m]) > 0.5))
mism_a = sorted(m for m, f in mism.items() if f != ["cx/pallet"] or abs(CAT[m]["paletizacao_caixas"] - CXPAL_A[m]) > 0.5)
check("DERIVED_MISMATCH: SKUs do JSON com valor ≠ dado mestre adotado (SAP + Q12)", ["12142501", "12144012", "12147801", "20908885"], mism_a)
# causa-raiz: máximo cru de todas as linhas de UM, inclusive a linha KG sintética
explain = sum(abs(max(x["net_raw"] for x in ROWS[m].values()) - INT[m]["peso_liquido_max"]) < 1e-9 for m in INT)
check("peso_liquido_max == máximo cru entre todas as UM (SKUs explicados)", 10, explain)
inconsistent = sorted(m for m in INT if abs(INT[m]["paletizacao_umb"] / ROWS[m]["CX"]["cnt"] - INT[m]["paletizacao_cx"]) > 0.5)
check("paletizacao_umb ÷ un/cx ≠ paletizacao_cx (SKUs)", ["12034103", "12037417", "12142501", "12144012", "20908885"], inconsistent)
names = {m for m in CAT if " ".join(CAT[m]["nome"].upper().split()) != " ".join(CX and next(r[3] for r in conf.iter_rows(min_row=2, values_only=True) if str(r[2]) == m).upper().split())}
check("nome do JSON ≠ Texto breve SAP (SKUs)", 5, len(names))
check("catálogo traz quantidade pedida embutida (itens)", 4, sum(c["caixas_pedidas"] > 0 for c in CAT.values()))
check("descrições ES/EN ausentes no JSON (SKUs)", ["12142146", "12147801"], sorted(m for m, c in CAT.items() if not c["descricao_es"] or not c["descricao_en"]))
check("AUX da OB sem ES/EN (SKUs)", ["12142146", "12147801"], sorted(m for m in SKUS if AUXLANG.get(m, set()) != {"PORTUGUÊS", "ESPAÑOL", "ENGLISH"}))
check("container.json: cubagem/peso (20', 40', carreta)", [(29.7, 25000.0), (60.84, 25000.0), (72.0, 25000.0)],
      [(VEH[k]["cubagem_m3"], VEH[k]["peso_bruto_kg"]) for k in ("container_20", "container_40", "carreta_padrao")])
check("OB: AA2 = peso/26 (dica de 26 posições)", round(20151.04 / 26, 2), round(ORDEM_AA2, 2) if ORDEM_AA2 else None)

section("4b) Resumo por família da OB, traduções e preço por kg")
ordem = ob["ORDEM"]
fam_resumo, fam_item = ordem["W11"].value, ordem["H22"].value
check("OB: rótulo da família no resumo (ORDEM!W11) × na linha do filtro (ORDEM!H22)", ("FILTRO", "FILTRO DE PAPEL"), (fam_resumo, fam_item))
fil = next(i for i in pedido["itens"] if i["material"] == 12037417)
check("valor e peso líquido que somem do resumo por família (US$, kg)", (7370.0, 230.0), (round(fil["valor_total"], 2), round(fil["peso_liquido"], 1)))
en_ = CAT["12034103"]["descricao_en"]
es_ = CAT["12034103"]["descricao_es"]
check("12034103: descrição en diz 200G e es diz 20G (produto real: 16×10×20G)", (True, True), ("200G" in en_, "20G" in es_ and "200G" not in es_))
ppk = lambda p, kg: round(p / kg, 2)
check("US$/kg líquido do cappuccino 12034103 e 12144012 (faixa SAP)", (7.73, 12.10), (ppk(SAP_PRICE["12034103"], 3.2), ppk(SAP_PRICE["12144012"], 0.8)))
check("US$/kg do 12034109 com o valor da OB (36,96) e com a estimativa antiga (100,38)", (7.70, 20.91), (ppk(36.96, 4.8), ppk(100.38, 4.8)))
check("valores dos 3 SKUs sem registro nos JSON v2 (= aba TAB da OB)", (36.96, 29.48, 17.611), (CAT["12034109"]["preco_venda"], CAT["12037417"]["preco_venda"], CAT["20908885"]["preco_venda"]))

# ================================================================ 5) stick 12147801
section("5) Peso bruto da caixa do stick 12147801: 0,89 kg (SAP) ou 1,2 kg (OB)?")
r = ROWS["12147801"]
n_sch = r["CX"]["cnt"] / r["SCH"]["cnt"]
area = 2 * (CX["12147801"]["L"] * CX["12147801"]["W"] + CX["12147801"]["L"] * CX["12147801"]["H"] + CX["12147801"]["W"] * CX["12147801"]["H"]) / 1e6
sub = r["SCH"]["gross"]
carton = lambda g: g - n_sch * sub
check("sachês por caixa", 240, n_sch)
check("peso bruto de 1 sachê (kg)", 0.0023, round(sub, 4), tol=1e-6)
check("candidato A 0,89 kg: papelão implícito (kg)", 0.338, round(carton(0.89), 3))
check("candidato B 1,2 kg: papelão implícito (kg)", 0.648, round(carton(1.2), 3))
dens = {}
for m in ("12034103", "12144012", "12142501", "20908885"):
    c, rr = CX[m], ROWS[m]
    a = 2 * (c["L"] * c["W"] + c["L"] * c["H"] + c["W"] * c["H"]) / 1e6
    s = rr.get("SCH") or rr.get("UN")
    dens[m] = (c["gross"] - c["un"] / s["cnt"] * s["gross"]) / a
check("papelão por m² dos SKUs de embalagem parecida (faixa, kg/m²)", (1.08, 1.46), (round(min(dens.values()), 2), round(max(dens.values()), 2)))
check("A 0,89 kg → kg/m²", 1.25, round(carton(0.89) / area, 2))
check("B 1,2 kg → kg/m²", 2.39, round(carton(1.2) / area, 2))

# ================================================================ 6) upsell (golden 2)
section("6) Upsell: pedido de 19 posições na carreta")
base = {"12031189": 2352, "12142501": 1400}
Tb = line_totals(base)
cap = VEH["carreta_padrao"]
ob_ = occupancy(Tb, "carreta_padrao", 28)
check("posições do pedido-base", 19, Tb["pos"])
check("posições livres", 9, 28 - Tb["pos"])
check("peso bruto (kg)", 14468.16, round(Tb["gross"], 2))
check("peso restante (kg)", 10531.84, round(cap["peso_bruto_kg"] - Tb["gross"], 2))
check("cubagem restante (m³)", 30.896, round(cap["cubagem_m3"] - Tb["m3_decl"], 3), tol=0.001)
check("ocupação por posições (%)", 67.9, round(ob_["pos_pct"], 1))
vb = sum(base[m] * SAP_PRICE[m] for m in base)
check("valor do pedido-base (US$)", 146396.32, round(vb, 2))
add = 9 * CXPAL["12031291"] * SAP_PRICE["12031291"]
check("9 pallets cheios de 12031291 (US$)", 90977.04, round(add, 2))
check("9 pallets de 12031189, já no pedido (US$)", 74707.92, round(9 * CXPAL["12031189"] * SAP_PRICE["12031189"], 2))
check("9 pallets de 12142501, já no pedido (US$)", 54331.20, round(9 * CXPAL["12142501"] * SAP_PRICE["12142501"], 2))
w_142501 = Tb["gross"] + 9 * CXPAL["12142501"] * CX["12142501"]["gross"]
check("peso depois de 9 pallets de 12142501 (% da carga útil)", 71.8, round(w_142501 / cap["peso_bruto_kg"] * 100, 1), tol=0.06)
w_189 = Tb["gross"] + 9 * CXPAL["12031189"] * CX["12031189"]["gross"]
check("peso depois de 9 pallets de 12031189 (% da carga útil)", 90.1, round(w_189 / cap["peso_bruto_kg"] * 100, 1), tol=0.06)
w_after = Tb["gross"] + 9 * CXPAL["12031291"] * CX["12031291"]["gross"]
check("peso depois de 9 pallets de 12031291 (% da carga útil)", 90.1, round(w_after / cap["peso_bruto_kg"] * 100, 1), tol=0.06)
heaviest = max(CXPAL_A[m] * CX[m]["gross"] for m in SKUS)
check("pallet mais pesado, com os valores adotados (kg)", 895.44, round(heaviest, 2))
check("9 pallets do mais pesado cabem no peso restante", True, 9 * heaviest <= cap["peso_bruto_kg"] - Tb["gross"])
bulkiest = max(CXPAL_A[m] * CX[m]["vol"] / 1000 for m in SKUS)
check("pallet de maior cubagem, com os valores adotados (m³; o 20908885 a 110 cx daria 2,788)", 2.230, round(bulkiest, 3), tol=0.001)
check("20908885 a 110 cx/pallet (SAP cru): cubagem do pallet (m³)", 2.788, round(CXPAL["20908885"] * CX["20908885"]["vol"] / 1000, 3), tol=0.001)
check("9 pallets do mais volumoso cabem na cubagem restante", True, 9 * bulkiest <= cap["cubagem_m3"] - Tb["m3_decl"])
lotes = {m: next((l for b, mm, d, l in PALET if mm == m and b == "LONT"), 1) for m in SKUS}
inviaveis = [m for m in SKUS if base.get(m, 0) / CXPAL_A[m] + 9 < lotes[m]]
check("SKUs em que 9 pallets não atendem o lote mínimo", [], inviaveis)
check("lote mínimo de 12142146 e 12142501 (pallets)", (5, 5), (next(l for b, m, d, l in PALET if m == "12142146"), next(l for b, m, d, l in PALET if m == "12142501")))

section("6b) As 20 divergências OB × SAP e a causa-raiz dos pesos no JSON")
OBTAB = {str(r[5]): r for r in ob["TAB"].iter_rows(min_row=3, values_only=True) if r[5]}
# Valores da OB, por SKU, resolvidos das abas:
#  - cx/pallet "PALETIZAÇÃO": un/pallet da filial LONT (coluna G) ÷ unidades da caixa do SAP; sem linha LONT, usa a de BHO
#  - lote mínimo "da OB": o que a ORDEM!Z busca por material (primeiro valor, filial BHO quando existe)
#  - lote mínimo "do SAP/LONT": a linha LONT; no 12037417 a de descrição "RED DES", igual à da Conferência
#  - peso bruto "da OB": peso bruto da unidade × unidades por caixa (ORDEM!M × ORDEM!J)
PALZ_LINES = [(r[1], str(r[2]), str(r[3]), r[6], r[7]) for r in ob["PALETIZAÇÃO"].iter_rows(min_row=3, values_only=True) if r[2]]


def palet_pick(m, filiais):
    for fil in filiais:
        rows = [x for x in PALZ_LINES if x[1] == m and x[0] == fil]
        if rows:
            return rows[0]
    return None


OBG = {str(r[5]): (r[12] or 0) * (r[9] or 0) for r in ob["ORDEM"].iter_rows(min_row=17, max_row=40, values_only=True) if r[5]}
div = {"cx/pallet OB x SAP": [], "cx/pallet aba PALETIZACAO x SAP": [], "preco": [], "sem preco SAP": [], "lote minimo": [], "peso bruto/cx": []}
for m in SKUS:
    c = CX[m]
    if abs(OBTAB[m][12] - CXPAL[m]) > 0.011:
        div["cx/pallet OB x SAP"].append(m)
    pl = palet_pick(m, ("LONT", "BHO"))
    if pl and abs(pl[3] / c["un"] - CXPAL[m]) > 0.011:
        div["cx/pallet aba PALETIZACAO x SAP"].append(m)
    if m in SAP_PRICE:
        if abs(OBTAB[m][8] - SAP_PRICE[m]) > 0.011:
            div["preco"].append(m)
    else:
        div["sem preco SAP"].append(m)
    ob_lote = (palet_pick(m, ("BHO", "LONT")) or (None, None, None, None, 1))[4]
    sap_rows = [x for x in PALZ_LINES if x[1] == m and x[0] == "LONT" and (m != "12037417" or "RED DES" in x[2])]
    sap_lote = sap_rows[0][4] if sap_rows else 1
    if abs(ob_lote - sap_lote) > 0.011:
        div["lote minimo"].append(m)
    if m in OBG and abs(OBG[m] - c["gross"]) > 0.05:
        div["peso bruto/cx"].append(m)
for k, v in div.items():
    print(f"  {k}: {len(v)} {v}")
check("divergências por tipo (cx/pallet OB, cx/pallet aba PALET, preço, sem preço, lote, peso bruto)", (2, 6, 1, 3, 2, 6), tuple(len(v) for v in div.values()))
check("total de divergências OB × SAP, recalculado das abas (README do piloto: 20)", 20, sum(len(v) for v in div.values()))
check("SKUs com peso bruto divergente na linha da OB (não são os mesmos do cadastro da caixa)", ["12031189", "12034103", "12037417", "12142146", "12147801", "20908885"], div["peso bruto/cx"])
print("  nota: as 20 divergências contam pares (campo × SKU); '6 de 10 SKUs com cadastro da caixa divergente' (seção 2b) conta SKUs, e os conjuntos diferem.")
check("peso_bruto_max == máximo cru entre todas as UM (SKUs explicados)", 10, sum(abs(max(x["gross_raw"] for x in ROWS[m].values()) - INT[m]["peso_bruto_max"]) < 1e-9 for m in INT))
check("preço do 12031189: a OB usa 49,4067 e o piloto guarda 49,41 (tolerância do piloto esconde)", (49.4067, 49.41), (round(OBTAB["12031189"][8], 4), SAP_PRICE["12031189"]))

section("6c) Estivada por piso (modelo ilustrativo) e volume pelas medidas")
stacks_area = sum(math.ceil(q / LAYERS[m][1]) * CX[m]["L"] * CX[m]["W"] / 1e6 for m, q in Q.items())
check("estivada: área de piso ocupada pelo pedido-teste, pilhas de 'altura' camadas (m²)", 32.43, round(stacks_area, 2))
check("estivada: % do piso da carreta se o piso for 28 × 1,2 m² (premissa)", 96.5, round(stacks_area / (28 * 1.2) * 100, 1))
check("estivada: com piso de 28 × 1,2 m² a carga de 84,1% de peso e 81,8% de cubagem ocupa o piso quase todo", True, stacks_area / (28 * 1.2) > 0.9)
check("estivada: % do piso do contêiner 40' se o piso for 20 × 1,2 m² (premissa)", 135.1, round(stacks_area / (20 * 1.2) * 100, 1))
check("estivada: % do piso do contêiner 20' se o piso for 10 × 1,2 m² (premissa)", 270.3, round(stacks_area / (10 * 1.2) * 100, 1))
check("estivada: piso da carreta, 40' e 20' pela premissa posições × 1,2 m² (m²)", (33.6, 24.0, 12.0), (round(28 * 1.2, 1), round(20 * 1.2, 1), round(10 * 1.2, 1)))
check("volume do pedido-teste pelas medidas C×L×A (m³)", 59.04, round(T["m3_dims"], 2))
check("carreta pelas medidas (%)", 82.0, round(T["m3_dims"] / 72 * 100, 1))
check("contêiner 40' pelas medidas (%)", 97.0, round(T["m3_dims"] / 60.84 * 100, 1))
check("upsell: valor de 9 pallets de 12031189, 12142501 e 12031291 (US$)", (74707.92, 54331.20, 90977.04),
      tuple(round(9 * CXPAL[m] * SAP_PRICE[m], 2) for m in ("12031189", "12142501", "12031291")))

# ================================================================ 7) pedido aprovado e restrição de país
section("7) Pedido aprovado do piloto (cliente demo Río Claro, contêiner 40') e restrição de país")


def pilot_round(x):  # Math.round do JavaScript (meio para cima)
    return math.floor(x * 100 + 0.5) / 100


# preço base da lista US; os 3 SKUs sem preço SAP usam o preço do site da 3corações (Q6), igual ao da aba TAB da OB
base_price = {m: (SAP_PRICE[m] if m in SAP_PRICE else TAB[m]) for m in SKUS}
AR = {m: pilot_round(base_price[m] * 1.06) for m in SKUS}  # lista AR simulada = US × 1,06 (cliente fictício do piloto)
q3 = {"12031291": 1176, "12034109": 140, "12142501": 1400, "12147801": 540}
STICK_GROSS = 1.2  # peso bruto do stick adotado pelo piloto (override da OB; ver seção 5 e Q9)
gross3 = sum(q * (STICK_GROSS if m == "12147801" else CX[m]["gross"]) for m, q in q3.items())
pos3 = sum(math.ceil(q / CXPAL[m] - EPS) for m, q in q3.items())  # variante do piloto: 12034109 a 28 cx/pallet (SAP cru)
m3_3 = sum(q * CX[m]["vol"] / 1000 for m, q in q3.items())
check("pedido 0002, variante do piloto (12034109 a 28 cx/pallet): posições no contêiner 40' (de 20)", 20, pos3)
check("pedido 0002: peso bruto (kg)", 9712.86, round(gross3, 2))
check("pedido 0002: cubagem (m³)", 33.15, round(m3_3, 2))
check("pedido 0002: valor com lista AR (US$)", 123761.08, round(sum(q * AR[m] for m, q in q3.items()), 2))
check("pedido 0002, variante do piloto: % posições / % cubagem / % peso", (100.0, 54.5, 38.9),
      (round(pos3 / 20 * 100, 1), round(m3_3 / VEH["container_40"]["cubagem_m3"] * 100, 1), round(gross3 / 25000 * 100, 1)))
pos3a = sum(math.ceil(q / CXPAL_A[m] - EPS) for m, q in q3.items())  # dados adotados: 12034109 a 98 (Q12)
status3a = "excesso" if max(pos3a / 20, m3_3 / 60.84, gross3 / 25000) > 1 else "completa" if max(pos3a / 20, m3_3 / 60.84, gross3 / 25000) >= 0.9 else "incompleta"
check("pedido 0002, dados adotados (98 cx/pallet, Q12): posições, % e status", (17, 85.0, "incompleta"), (pos3a, round(pos3a / 20 * 100, 1), status3a))
p109 = q3["12034109"] / CXPAL_A["12034109"]
check("pedido 0002, dados adotados: 12034109 em pallets e caixas para fechar o segundo", (1.43, 56),
      (round(p109, 2), math.ceil(CXPAL_A["12034109"] * (math.floor(p109 + EPS) + 1) - q3["12034109"] - EPS)))
check("pedido 0002, dados adotados: posições livres no contêiner 40'", 3, 20 - pos3a)
check("pedido 0002: com o bruto do stick em 0,89 kg (kg)", 9545.46, round(gross3 - 540 * (STICK_GROSS - 0.89), 2))
check("preço AR de 12142501 (US × 1,06)", 22.85, AR["12142501"])
check("1010 UN de 12031291 (20 un/cx) = caixas", 50.5, 1010 / CX["12031291"]["un"])
below_min = 560 / CXPAL["12142501"] < next(l for b, m, d, l in PALET if m == "12142501" and b == "LONT")
check("560 cx de 12142501 = 2 pallets, abaixo do lote mínimo de 5", True, below_min)
p885 = 110 / CXPAL_A["20908885"]
check("caso 4, dados adotados: 110 cx de 20908885 a 77 cx/pallet (pallets, posições, caixas para fechar)", (1.43, 2, 44),
      (round(p885, 2), math.ceil(p885 - EPS), math.ceil(CXPAL_A["20908885"] * (math.floor(p885 + EPS) + 1) - 110 - EPS)))
check("caso 4, variante do piloto: 110 cx de 20908885 a 110 cx/pallet (pallets)", 1.0, 110 / CXPAL["20908885"])

# ================================================================ 8) confronto com o motor do próprio piloto
section("8) Confronto com o motor do piloto (opcional)")
import os
import subprocess

pilot_html = os.environ.get("PILOTO_INDEX_HTML") or (sys.argv[1] if len(sys.argv) > 1 else None)
if not pilot_html or not Path(pilot_html).exists():
    print("  [pulado] informe o index.html do piloto (commit 636ac8d) em PILOTO_INDEX_HTML ou como argumento para comparar.")
else:
    harness = Path(__file__).with_name("motor_piloto_harness.js")
    res = subprocess.run(["node", str(harness), pilot_html], capture_output=True, text=True, encoding="utf-8")
    if res.returncode != 0:
        check("harness do piloto executou", 0, res.returncode)
        print(res.stderr[:400])
    else:
        P_ = json.loads(res.stdout)
        c1 = P_["caso1_carreta"]
        check("piloto × Python: peso bruto do caso 1 (kg)", round(T["gross"], 2), c1["gross"])
        check("piloto × Python: peso líquido do caso 1 (kg)", round(T["net"], 2), c1["net"])
        check("piloto × Python: posições do caso 1", T["pos"], c1["pos"])
        check("piloto × Python: valor do caso 1 (US$)", round(valor, 2), c1["value"])
        check("piloto × Python: % de posições da carreta no caso 1", round(o["pos_pct"], 1), c1["pct"]["pos"])
        check("piloto × Python: pallet misto no caso 1 (posições)", 29, P_["caso1_carreta_misto"]["pos"])
        check("piloto × Python: contêiner 40' no caso 1 (% de posições, % de cubagem)", (150.0, 96.8), (P_["caso1_c40"]["pct"]["pos"], P_["caso1_c40"]["pct"]["m3"]))
        check("piloto: estivada na carreta no caso 1 é 'incompleta' (peso e cubagem)", ("partial", "kg"), (P_["caso1_estivada_carreta"]["status"], P_["caso1_estivada_carreta"]["bind"]))
        check("piloto × Python: caso 2, posições e peso (kg)", (Tb["pos"], round(Tb["gross"], 2)), (P_["caso2_upsell"]["pos"], P_["caso2_upsell"]["gross"]))
        check("piloto × Python: caso 3, peso e valor", (round(gross3, 2), 123761.08), (P_["caso3_pedido0002"]["gross"], P_["caso3_pedido0002"]["value"]))
        check("piloto × Python: caso 3, posições na variante do piloto (12034109 a 28 cx/pallet)", pos3, P_["caso3_pedido0002"]["pos"])
        check("piloto: pallets do 12034109 no caso 3 e do 20908885 no caso 4 (dados do piloto, antes da Q12)", (5.0, 1.0),
              (P_["caso3_pedido0002"]["pal"]["12034109"], P_["caso4_pais"]["pal"]["20908885"]))
        check("piloto: caso 4 bloqueios", ["12142501:belowMin", "20908885:blocked"], sorted(P_["caso4_pais"]["issues"]))
        check("piloto: posições dos veículos", {"carreta": 28, "c40": 20, "c20": 10}, {k: v["pos"] for k, v in P_["veiculos"].items()})

# ================================================================ resumo
ok = sum(1 for c in CHECKS if c[3])
print(f"\n== Resumo: {ok} de {len(CHECKS)} verificações conferem")
for d, e, g, k in CHECKS:
    if not k:
        print(f"   DIVERGE: {d} | esperado {e} | obtido {g}")
sys.exit(0 if ok == len(CHECKS) else 1)
