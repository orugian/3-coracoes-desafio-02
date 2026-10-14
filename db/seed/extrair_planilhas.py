#!/usr/bin/env python3
"""Extrai as planilhas do desafio para db/seed/fontes/*.json (docs/specs/03-seed.md, decisão D6).

Uso (na raiz do repositório; normalmente via `npm run extract`):
    python db/seed/extrair_planilhas.py             # grava as fontes
    python db/seed/extrair_planilhas.py --verificar # só confere que nenhuma senha da OB está nas fontes gravadas

Mesma leitura do docs/arquitetura/verificacao/verificar_numeros.py (mesmas colunas e normalização). Cada arquivo traz
{ source, file, sha256, extractedAt, rows }. O sha256 é dos bytes do arquivo de origem com \\r\\n trocado por \\n.
extractedAt só muda quando o conteúdo muda, para regenerar sem diferença no git.

Segurança: a aba DADOS CLIENTES tem a senha do cliente na coluna [2] e embutida na [0] (CHAVE = código + senha).
O extrator nunca lê essas colunas para a saída (lista fechada de colunas) e, antes de gravar, procura cada valor como
substring no JSON serializado; se achar, aborta sem imprimir o valor.
"""
import datetime as dt
import glob
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "pre-solucao" / "pedidos-sem-fronteiras"
OUT = ROOT / "db" / "seed" / "fontes"
SKUS_RE_LEN = 8


class SecretLeak(Exception):
    """Uma senha da DADOS CLIENTES apareceria numa fonte gerada. A mensagem nunca traz o valor."""


def first(pattern):
    found = sorted(glob.glob(str(SRC / pattern)))
    if len(found) != 1:
        raise SystemExit(f"esperava 1 arquivo para {pattern} em {SRC}, achei {len(found)}")
    return Path(found[0])


def sha256_normalized(path):
    return hashlib.sha256(path.read_bytes().replace(b"\r\n", b"\n")).hexdigest()


def rel(path):
    return path.resolve().relative_to(ROOT).as_posix()


def code(value):
    """Material, cliente ou centro como texto ('12031189'), venha como int, float inteiro ou texto."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def iso_date(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raise SystemExit(f"data inesperada: {value!r}")


def num(value):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise SystemExit(f"número inesperado: {value!r}")
    return value


# ---------------------------------------------------------------- SAP
def extract_conferencia(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows, seen = [], set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[2] is None:
            continue
        material, uom = code(r[2]), r[5]
        if (material, uom) in seen:
            raise SystemExit(f"Conferência: UM repetida {material}/{uom}")
        seen.add((material, uom))
        if r[12] not in (None, "MM") or r[14] not in (None, "CD3") or r[17] not in ("KG", "G"):
            raise SystemExit(f"Conferência: unidade inesperada em {material}/{uom}: {r[12]}, {r[14]}, {r[17]}")
        rows.append({
            "centro": code(r[0]), "material": material, "texto": r[3], "denominador": num(r[4]), "um": uom,
            "contador": num(r[6]), "ean": code(r[7]), "ctgEan": r[8],
            "comprimentoMm": num(r[9]) if r[12] else None, "larguraMm": num(r[10]) if r[12] else None,
            "alturaMm": num(r[11]) if r[12] else None, "volumeDm3": num(r[13]) if r[14] else None,
            "liquido": num(r[15]), "bruto": num(r[16]), "unidadePeso": r[17],
            "validadeDias": num(r[20]), "hierarquia": code(r[21]), "lastro": num(r[26]), "camadas": num(r[27]),
        })
    return rows


def extract_materiais(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    return [{"material": code(r[0]), "descricao": r[1], "centro": code(r[4]), "lastro": num(r[5]), "altura": num(r[6]),
             "embalagem": r[7]} for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def extract_registros(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    return [{"canal": code(r[0]), "centro": code(r[2]), "descCentro": r[3], "lista": r[6], "descLista": r[7],
             "material": code(r[8]), "um": r[12], "validoDesde": iso_date(r[13]), "validoAte": iso_date(r[14]),
             "precoSugerido": num(r[16]), "moeda": r[18], "exclusao": num(r[19])}
            for r in ws.iter_rows(min_row=2, values_only=True) if r[8] is not None]


# ---------------------------------------------------------------- Order Basis
def is_material(value):
    c = code(value)
    return c is not None and c.isdigit() and len(c) == SKUS_RE_LEN


def client_rows(ob):
    """Linhas de cliente da DADOS CLIENTES: todas as que vêm depois do cabeçalho ([1] == "Cliente") e têm código.
    Não depende do tipo do código, para um cliente com código em texto não escapar da autoverificação."""
    rows = list(ob["DADOS CLIENTES"].iter_rows(values_only=True))
    header = next((i for i, r in enumerate(rows) if len(r) > 1 and isinstance(r[1], str) and r[1].strip().lower() == "cliente"), None)
    if header is None:
        raise SystemExit("DADOS CLIENTES sem a linha de cabeçalho (coluna [1] = 'Cliente')")
    return [r for r in rows[header + 1:] if len(r) > 1 and code(r[1])]


def client_secrets(ob):
    """Valores de senha da DADOS CLIENTES (colunas [2] e [0]) das linhas de cliente. Nunca impressos."""
    secrets = set()
    for r in client_rows(ob):
        for value in (r[2] if len(r) > 2 else None, r[0]):
            if value is not None and code(value):
                secrets.add(code(value))
    return secrets


def extract_ob(path):
    ob = openpyxl.load_workbook(path, data_only=True, keep_vba=True, read_only=True)
    clientes = []
    for r in client_rows(ob):
        if len(r) < 14:
            raise SystemExit(f"DADOS CLIENTES: linha de cliente com {len(r)} colunas (esperado 14)")
        # Lista fechada de colunas: [0] (CHAVE) e [2] (senha) nunca entram.
        clientes.append({"codigo": code(r[1]), "resumo": r[3], "lista": r[4], "nome": r[5], "condPagamentoCodigo": code(r[6]),
                         "condPagamentoTexto": r[7], "incoterm": r[8], "ativo": r[9], "idioma": r[10], "moeda": r[11],
                         "filial": r[13]})
    tab = [{"cliente": code(r[2]), "org": r[3], "lista": r[4], "material": code(r[5]), "descricao": r[6], "um": r[7],
            "preco": num(r[8]), "moeda": r[9], "umb": r[10], "paletizacaoUmb": num(r[11]), "paletizacaoCx": num(r[12])}
           for r in ob["TAB"].iter_rows(min_row=3, values_only=True) if is_material(r[5])]
    paletizacao = [{"org": r[1], "material": code(r[2]), "descricao": r[3], "um": r[4], "unidadesPallet": num(r[6]),
                    "loteMinimo": num(r[7])}
                   for r in ob["PALETIZAÇÃO"].iter_rows(min_row=3, values_only=True) if is_material(r[2])]
    aux = [{"idioma": r[2], "material": code(r[3]), "descricao": r[4], "familia": r[5]}
           for r in ob["AUX"].iter_rows(min_row=1, max_row=60, values_only=True)
           if len(r) > 5 and r[2] in ("PORTUGUÊS", "ESPAÑOL", "ENGLISH") and is_material(r[3])]
    ordem = ob["ORDEM"]
    cell = lambda a: ordem[a].value
    itens = [{"material": code(r[5]), "descricao": r[6], "familia": r[7], "preco": num(r[8]), "unidadesCaixa": num(r[9]),
              "brutoUnidade": num(r[12]), "liquidoUnidade": num(r[13]), "cxPallet": num(r[18]), "caixas": num(r[19]),
              "loteMinimo": num(r[25]), "validado": r[26]}
             for r in ordem.iter_rows(min_row=5, max_row=60, values_only=True) if len(r) > 26 and is_material(r[5])]
    totais = {"X2_pesoBruto": num(cell("X2")), "W2_pesoLiquido": num(cell("W2")), "Y2_valor": num(cell("Y2")),
              "X4_cubagem": num(cell("X4")), "Y4_pallets": num(cell("Y4")), "W4_embarque": cell("W4")}
    rows = {"clientes": clientes, "tab": tab, "paletizacao": paletizacao, "aux": aux, "ordem": {"totais": totais, "itens": itens}}
    return rows, client_secrets(ob)


# ---------------------------------------------------------------- gravação
def assert_no_secrets(text, secrets):
    # Procura o valor cru e a forma escapada em JSON (uma senha com aspas, barra ou acento viraria \", \\ ou \uXXXX).
    for i, secret in enumerate(sorted(secrets)):
        forms = {secret, json.dumps(secret)[1:-1], json.dumps(secret, ensure_ascii=False)[1:-1]}
        if any(f and f in text for f in forms):
            raise SecretLeak(f"valor de senha nº {i + 1} da aba DADOS CLIENTES apareceria na fonte gerada")


def envelope(source, src_path, rows, out_path):
    sha = sha256_normalized(src_path)
    extracted_at = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    if out_path.exists():
        old = json.loads(out_path.read_text(encoding="utf-8"))
        if old.get("sha256") == sha and old.get("rows") == rows:
            extracted_at = old.get("extractedAt", extracted_at)
    return {"source": source, "file": rel(src_path), "sha256": sha, "extractedAt": extracted_at, "rows": rows}


def serialize(doc):
    return json.dumps(doc, ensure_ascii=False, indent=2) + "\n"


def extract():
    conf, mats, regs = first("Confer*.xlsx"), first("Materiais*.xlsx"), first("Registros*.xlsx")
    ob_path = SRC / "ORDEN BASIS TESTE.xlsm"
    ob_rows, secrets = extract_ob(ob_path)
    if not secrets:
        raise SystemExit("DADOS CLIENTES sem linha de cliente: a autoverificação de senha não teria o que procurar")
    docs = {
        "sap_conferencia.json": ("SAP_CONFERENCIA", conf, extract_conferencia(conf)),
        "sap_materiais.json": ("SAP_MATERIAIS", mats, extract_materiais(mats)),
        "sap_registros.json": ("SAP_REGISTROS", regs, extract_registros(regs)),
        "ob.json": ("ORDER_BASIS", ob_path, ob_rows),
    }
    texts = {}
    for name, (source, src, rows) in docs.items():
        texts[name] = serialize(envelope(source, src, rows, OUT / name))
        assert_no_secrets(texts[name], secrets)  # antes de gravar qualquer arquivo
    OUT.mkdir(parents=True, exist_ok=True)
    for name, text in texts.items():
        with open(OUT / name, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)
        print(f"  {rel(OUT / name)}")
    print(f"fontes do SAP e da OB gravadas; autoverificação: {len(secrets)} valores de senha procurados, nenhum encontrado")


def verify():
    """Confere as fontes já gravadas (todas as de db/seed/fontes/) e o próprio detector."""
    ob = openpyxl.load_workbook(SRC / "ORDEN BASIS TESTE.xlsm", data_only=True, keep_vba=True, read_only=True)
    secrets = client_secrets(ob)
    if not secrets:
        raise SystemExit("DADOS CLIENTES sem linha de cliente")
    planted = next(iter(secrets))
    try:
        assert_no_secrets('{"x": "a' + planted + 'b"}', secrets)
        raise SystemExit("detector não achou um valor plantado")
    except SecretLeak:
        pass
    files = sorted(OUT.glob("*.json"))
    for f in files:
        assert_no_secrets(f.read_text(encoding="utf-8"), secrets)
    print(f"autoverificação ok: {len(secrets)} valores de senha procurados em {len(files)} fontes, nenhum encontrado; detector ok")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        verify() if "--verificar" in sys.argv[1:] else extract()
    except SecretLeak as e:
        print(f"ABORTADO: {e}. Nada foi gravado.", file=sys.stderr)
        sys.exit(3)
